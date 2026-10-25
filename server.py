"""Skill Passport API.

Run locally with: uvicorn server:app --reload --port 8000
"""
from contextlib import asynccontextmanager
import csv
import hashlib
import hmac
import io
import json
import os
import re
import secrets
import threading
import uuid
import zipfile
from collections import Counter
from pathlib import Path
from typing import Literal

from fastapi import FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from pypdf import PdfReader
from pypdf.errors import FileNotDecryptedError, PdfReadError
from starlette.concurrency import run_in_threadpool

ROOT = Path(__file__).resolve().parent
DATABASE_URL = os.getenv("DATABASE_URL", "").replace("postgres://", "postgresql://", 1)
DB_FILE = Path(os.getenv("SKILL_PASSPORT_DB_FILE", ROOT / "db.json"))
MAX_RESUME_BYTES = 10 * 1024 * 1024
SKILLS = [
    "python", "java", "javascript", "c++", "sql", "html", "css", "react",
    "machine learning", "deep learning", "artificial intelligence", "data science",
    "data analysis", "pandas", "numpy", "tensorflow", "pytorch", "computer vision",
    "opencv", "nlp", "natural language processing", "statistics", "excel",
    "power bi", "tableau", "git", "github", "docker", "linux", "cybersecurity",
    "networking", "figma", "ui design", "ux design",
]


def clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9\s+#]", " ", text.lower())).strip()


def extract_skills(text: str) -> list[str]:
    lowered = text.lower()
    return [skill for skill in SKILLS if re.search(r"\b" + re.escape(skill) + r"\b", lowered)]


def text_vector(text: str) -> Counter[str]:
    return Counter(re.findall(r"[a-z0-9+#]+", clean_text(text)))


def cosine_similarity(left: str, right: str) -> float:
    left_vector, right_vector = text_vector(left), text_vector(right)
    dot = sum(count * right_vector.get(word, 0) for word, count in left_vector.items())
    left_norm = sum(count * count for count in left_vector.values())
    right_norm = sum(count * count for count in right_vector.values())
    denominator = (left_norm * right_norm) ** 0.5
    return dot / denominator if denominator else 0.0


def opportunity_text(opportunity: dict) -> str:
    return " ".join([
        opportunity.get("title", ""),
        opportunity.get("description", ""),
        *opportunity.get("skills", []),
    ])


def profile_text(profile: dict) -> str:
    parts = list(profile.get("skills", []))
    for project in profile.get("projects", []):
        parts.extend([
            project.get("title", ""),
            project.get("description", ""),
            project.get("tech", ""),
        ])
    parts.extend(item.get("title", "") for item in profile.get("achievements", []))
    parts.append(profile.get("resume_text", ""))
    return " ".join(parts)


def extract_docx_text(raw: bytes) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(io.BytesIO(raw))
    except (PackageNotFoundError, zipfile.BadZipFile) as error:
        raise HTTPException(status_code=400, detail="The uploaded resume could not be read.") from error
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join([*paragraphs, *table_cells])


def empty_state() -> dict:
    opportunities = []
    csv_path = ROOT / "data" / "opportunities.csv"
    if csv_path.exists():
        with csv_path.open(encoding="utf-8", newline="") as source:
            for index, row in enumerate(csv.DictReader(source), start=1):
                description = row.get("description", "")
                opportunities.append({
                    "id": f"o{index}",
                    "title": row.get("title", ""),
                    "org": row.get("organization", "Campus"),
                    "type": row.get("type", "Project"),
                    "description": description,
                    "skills": extract_skills(description),
                    "deadline": "",
                    "contact": "",
                    "creator": None,
                })
    return {
        "users": {},
        "profiles": {},
        "opportunities": opportunities,
        "applications": [],
        "saved": [],
        "sessions": {},
    }


def load_state() -> dict:
    if DATABASE_URL:
        import psycopg
        from psycopg.types.json import Jsonb

        with psycopg.connect(DATABASE_URL) as connection:
            connection.execute(
                "CREATE TABLE IF NOT EXISTS skill_passport_state "
                "(id SMALLINT PRIMARY KEY CHECK (id = 1), data JSONB NOT NULL)"
            )
            row = connection.execute(
                "SELECT data FROM skill_passport_state WHERE id = 1"
            ).fetchone()
            if row:
                return row[0]
            state = empty_state()
            connection.execute(
                "INSERT INTO skill_passport_state (id, data) VALUES (1, %s)",
                (Jsonb(state),),
            )
            return state
    if DB_FILE.exists():
        with DB_FILE.open(encoding="utf-8") as source:
            return json.load(source)
    state = empty_state()
    save_state(state)
    return state


def save_state(state: dict) -> None:
    if DATABASE_URL:
        import psycopg
        from psycopg.types.json import Jsonb

        with psycopg.connect(DATABASE_URL) as connection:
            connection.execute(
                "INSERT INTO skill_passport_state (id, data) VALUES (1, %s) "
                "ON CONFLICT (id) DO UPDATE SET data = EXCLUDED.data",
                (Jsonb(state),),
            )
        return
    temporary_path = DB_FILE.with_suffix(DB_FILE.suffix + ".tmp")
    with temporary_path.open("w", encoding="utf-8") as target:
        json.dump(state, target, indent=1)
    temporary_path.replace(DB_FILE)


def hash_password(password: str, salt: bytes | None = None) -> tuple[str, str]:
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.scrypt(password.encode("utf-8"), salt=salt, n=2**14, r=8, p=1)
    return salt.hex(), digest.hex()


def public_user(user: dict) -> dict:
    return {key: value for key, value in user.items() if key not in {"pw", "pw_salt"}}


def current_user(authorization: str | None) -> dict:
    token = (authorization or "").removeprefix("Bearer ").strip()
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    user_id = db["sessions"].get(token_hash)
    if not user_id or user_id not in db["users"]:
        raise HTTPException(status_code=401, detail="Please sign in again.")
    return db["users"][user_id]


def ranked(profile: dict) -> list[dict]:
    text = profile_text(profile)
    opportunities = [
        {**item, "similarity": cosine_similarity(text, opportunity_text(item))}
        for item in db["opportunities"]
    ]
    return sorted(opportunities, key=lambda item: -item["similarity"])


def seed_state_if_needed() -> None:
    global db
    if os.getenv("REQUIRE_DATABASE_URL", "").lower() == "true" and not DATABASE_URL:
        raise RuntimeError("DATABASE_URL must be set for the hosted service.")
    db.clear()
    db.update(load_state())


db: dict = {}
request_lock = threading.Lock()


@asynccontextmanager
async def lifespan(_app):
    await run_in_threadpool(seed_state_if_needed)
    yield


app = FastAPI(lifespan=lifespan)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "https://dharshini-hemachandran.github.io",
        "http://localhost:8000",
        "http://127.0.0.1:8000",
        *([os.environ["RENDER_EXTERNAL_URL"]] if os.getenv("RENDER_EXTERNAL_URL") else []),
    ],
    allow_methods=["GET", "POST", "PUT", "DELETE"],
    allow_headers=["Authorization", "Content-Type"],
)


@app.middleware("http")
async def refresh_shared_state(request, call_next):
    await run_in_threadpool(request_lock.acquire)
    try:
        latest_state = await run_in_threadpool(load_state)
        db.clear()
        db.update(latest_state)
        return await call_next(request)
    finally:
        request_lock.release()


class Signup(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=10, max_length=128)
    role: Literal["student", "creator"] = "student"
    field: str = ""
    year: str = ""
    org: str = ""
    org_type: str = ""
    position: str = ""


class Login(BaseModel):
    email: str = Field(min_length=3, max_length=254)
    password: str = Field(min_length=1, max_length=128)


@app.get("/api/health")
def health() -> dict:
    return {"status": "ok"}


@app.post("/api/signup")
def signup(details: Signup) -> dict:
    email = details.email.strip().lower()
    if not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        raise HTTPException(status_code=422, detail="Enter a valid email address.")
    if any(user["email"] == email for user in db["users"].values()):
        raise HTTPException(status_code=409, detail="Email already registered.")
    user_id = uuid.uuid4().hex
    token = secrets.token_urlsafe(32)
    salt, password_hash = hash_password(details.password)
    user = details.model_dump(exclude={"password"})
    user.update({"id": user_id, "email": email, "pw": password_hash, "pw_salt": salt})
    db["users"][user_id] = user
    db["profiles"][user_id] = {
        "skills": [],
        "projects": [],
        "achievements": [],
        "links": [],
        "evidence": [],
        "resume_text": "",
    }
    db["sessions"][hashlib.sha256(token.encode("utf-8")).hexdigest()] = user_id
    save_state(db)
    return {"token": token, "user": public_user(user)}


@app.post("/api/login")
def login(details: Login) -> dict:
    email = details.email.strip().lower()
    user = next(
        (item for item in db["users"].values() if item["email"] == email),
        None,
    )
    if user is None:
        raise HTTPException(status_code=401, detail="Invalid email or password.")
    salt, expected_hash = bytes.fromhex(user["pw_salt"]), user["pw"]
    _, actual_hash = hash_password(details.password, salt)
    if not hmac.compare_digest(actual_hash, expected_hash):
        raise HTTPException(status_code=401, detail="Invalid email or password.")
    token = secrets.token_urlsafe(32)
    db["sessions"][hashlib.sha256(token.encode("utf-8")).hexdigest()] = user["id"]
    save_state(db)
    return {"token": token, "user": public_user(user)}


@app.post("/api/forgot-password")
def forgot_password() -> dict:
    raise HTTPException(status_code=501, detail="Password recovery is not configured yet.")


@app.post("/api/logout")
def logout(authorization: str | None = Header(default=None)) -> dict:
    token = (authorization or "").removeprefix("Bearer ").strip()
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    if db["sessions"].pop(token_hash, None) is not None:
        save_state(db)
    return {"ok": True}


@app.get("/api/profile")
def get_profile(authorization: str | None = Header(default=None)) -> dict:
    user = current_user(authorization)
    return {"user": public_user(user), "profile": db["profiles"][user["id"]]}


@app.put("/api/profile")
def put_profile(profile: dict, authorization: str | None = Header(default=None)) -> dict:
    user = current_user(authorization)
    db["profiles"][user["id"]] = profile
    save_state(db)
    return profile


@app.put("/api/user")
def update_user(details: dict, authorization: str | None = Header(default=None)) -> dict:
    user = current_user(authorization)
    for key in ("name", "field"):
        if key in details:
            user[key] = details[key]
    save_state(db)
    return public_user(user)


@app.post("/api/analyze-resume")
async def analyze_resume(
    file: UploadFile = File(...),
    authorization: str | None = Header(default=None),
) -> dict:
    current_user(authorization)
    filename = file.filename or ""
    extension = Path(filename).suffix.lower()
    if extension not in {".txt", ".pdf", ".docx"}:
        raise HTTPException(
            status_code=415,
            detail="Supported resume formats are .txt, .pdf, and .docx.",
        )
    raw = await file.read(MAX_RESUME_BYTES + 1)
    if len(raw) > MAX_RESUME_BYTES:
        raise HTTPException(status_code=413, detail="Resume files must be 10 MB or smaller.")
    try:
        if extension == ".pdf":
            text = "\n".join(
                page.extract_text() or ""
                for page in PdfReader(io.BytesIO(raw)).pages
            )
        elif extension == ".docx":
            text = extract_docx_text(raw)
        else:
            text = raw.decode("utf-8")
    except (
        FileNotDecryptedError,
        PdfReadError,
        UnicodeDecodeError,
        ValueError,
        OSError,
        zipfile.BadZipFile,
    ) as error:
        raise HTTPException(status_code=400, detail="The uploaded resume could not be read.") from error
    if not text.strip():
        raise HTTPException(
            status_code=422,
            detail="No readable text was found. Scanned PDFs need OCR.",
        )
    return {"skills": extract_skills(text), "text": text}


@app.post("/api/match-opportunities")
def match_opportunities(profile: dict) -> list[dict]:
    return ranked(profile)


@app.get("/api/opportunities")
def list_opportunities() -> list[dict]:
    return db["opportunities"]


@app.post("/api/opportunities")
def create_opportunity(
    opportunity: dict,
    authorization: str | None = Header(default=None),
) -> dict:
    user = current_user(authorization)
    if user["role"] != "creator":
        raise HTTPException(status_code=403, detail="Only opportunity creators can post opportunities.")
    if not opportunity.get("title") or not opportunity.get("description"):
        raise HTTPException(status_code=400, detail="Title and description required.")
    item = {
        **opportunity,
        "id": "o" + uuid.uuid4().hex[:8],
        "creator": user["id"],
    }
    item["skills"] = item.get("skills") or extract_skills(item["description"])
    db["opportunities"].append(item)
    save_state(db)
    return item


def find_opportunity(opportunity_id: str) -> dict:
    item = next(
        (opportunity for opportunity in db["opportunities"] if opportunity["id"] == opportunity_id),
        None,
    )
    if item is None:
        raise HTTPException(status_code=404, detail="Opportunity not found.")
    return item


@app.get("/api/opportunities/{opportunity_id}")
def get_opportunity(opportunity_id: str) -> dict:
    return find_opportunity(opportunity_id)


@app.put("/api/opportunities/{opportunity_id}")
def update_opportunity(
    opportunity_id: str,
    details: dict,
    authorization: str | None = Header(default=None),
) -> dict:
    user = current_user(authorization)
    item = find_opportunity(opportunity_id)
    if user["role"] != "creator" or item["creator"] != user["id"]:
        raise HTTPException(status_code=403, detail="You can only edit your own opportunities.")
    item.update({key: value for key, value in details.items() if key not in {"id", "creator"}})
    save_state(db)
    return item


@app.delete("/api/opportunities/{opportunity_id}")
def delete_opportunity(
    opportunity_id: str,
    authorization: str | None = Header(default=None),
) -> dict:
    user = current_user(authorization)
    item = find_opportunity(opportunity_id)
    if user["role"] != "creator" or item["creator"] != user["id"]:
        raise HTTPException(status_code=403, detail="You can only delete your own opportunities.")
    db["opportunities"].remove(item)
    save_state(db)
    return {"ok": True}


@app.post("/api/opportunities/{opportunity_id}/apply")
def apply_to_opportunity(
    opportunity_id: str,
    authorization: str | None = Header(default=None),
) -> dict:
    user = current_user(authorization)
    if user["role"] != "student":
        raise HTTPException(status_code=403, detail="Only students can apply to opportunities.")
    find_opportunity(opportunity_id)
    application = {"opp": opportunity_id, "user": user["id"]}
    if application not in db["applications"]:
        db["applications"].append(application)
        save_state(db)
    return {"status": "APPLICATION SUBMITTED"}


@app.post("/api/opportunities/{opportunity_id}/save")
def toggle_saved_opportunity(
    opportunity_id: str,
    authorization: str | None = Header(default=None),
) -> dict:
    user = current_user(authorization)
    if user["role"] != "student":
        raise HTTPException(status_code=403, detail="Only students can save opportunities.")
    find_opportunity(opportunity_id)
    saved = {"opp": opportunity_id, "user": user["id"]}
    if saved in db["saved"]:
        db["saved"].remove(saved)
        is_saved = False
    else:
        db["saved"].append(saved)
        is_saved = True
    save_state(db)
    return {"saved": is_saved}


@app.get("/api/my-status")
def my_status(authorization: str | None = Header(default=None)) -> dict:
    user = current_user(authorization)
    return {
        "applied": [
            item["opp"] for item in db["applications"] if item["user"] == user["id"]
        ],
        "saved": [item["opp"] for item in db["saved"] if item["user"] == user["id"]],
    }


@app.get("/api/opportunities/{opportunity_id}/candidates")
def candidates(
    opportunity_id: str,
    authorization: str | None = Header(default=None),
) -> list[dict]:
    user = current_user(authorization)
    opportunity = find_opportunity(opportunity_id)
    if user["role"] != "creator" or opportunity["creator"] != user["id"]:
        raise HTTPException(status_code=403, detail="You can only view candidates for your own opportunities.")
    ranked_candidates = []
    for user_id, profile in db["profiles"].items():
        candidate_user = db["users"][user_id]
        if candidate_user["role"] != "student":
            continue
        visible_profile = {key: value for key, value in profile.items() if key != "resume_text"}
        ranked_candidates.append({
            "id": user_id,
            "name": candidate_user["name"],
            "field": candidate_user.get("field", ""),
            "profile": visible_profile,
            "similarity": cosine_similarity(
                opportunity_text(opportunity),
                profile_text(profile),
            ),
        })
    return sorted(ranked_candidates, key=lambda item: -item["similarity"])


@app.get("/")
def index() -> FileResponse:
    return FileResponse(ROOT / "static" / "index.html")
